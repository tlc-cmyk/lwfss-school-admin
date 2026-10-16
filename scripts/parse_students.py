#!/usr/bin/env python3
"""Parse class rosters + subject groups → JSON + SQL seeds for Supabase."""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl
import xlrd

ATT = Path(
    "/cursor/stores/store-b946d782-e9c8-484f-a981-79a6c848f901/"
    "grok-bot/agents/7cbf6b5a-cad3-4460-8974-eb45b7fc8aef/attachments"
)
OUT = Path("/workspace/school-admin/scripts/_student_parse")
YEAR = "2026-2027"

# staff class_teacher → short_name (from Supabase; lowercased)
STAFF_CLASS_TEACHERS = {
    "1A": ["wcy", "ayws"],
    "1B": ["hsm", "nhh"],
    "1C": ["yth", "nckk"],
    "1D": ["lpc", "lttj"],
    "2A": ["hsy", "kmw"],
    "2B": ["syt", "wpc"],
    "2C": ["yshc", "kkh"],
    "2D": ["cyt", "wsfa", "wwkf"],
    "3A": ["pyl", "hcs"],
    "3B": ["wsf", "ltc"],
    "3C": ["lty", "ccy"],
    "3D": ["tsc", "hv"],
    "4A": ["lwy", "cmwe"],
    "4B": ["lchj", "kns"],
    "4C": ["ckyt", "ncy"],
    "4D": ["tlp", "wks"],
    "5A": ["hty", "lkh"],
    "5B": ["lyf", "nyw"],
    "5C": ["cwl", "scl"],
    "5D": ["ckyi", "fwk"],
    "6A": ["lna", "cyk"],
    "6B": ["ccc", "lpl"],
    "6C": ["ckhc", "wkyt"],
    "6D": ["kwy", "ltt"],
}

SKIP_SHEETS = {
    "All",
    "工作表1",
    "數學分布",
    "APL 全部學生",
    "S3升S4 選修科名單",
    "工作紙",
}


def sql_str(s) -> str:
    if s is None:
        return "NULL"
    t = str(s).replace("\xa0", " ").strip()
    t = t.replace("'", "''")
    return f"'{t}'"


def parse_students_xlsx(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    students = []
    for name in wb.sheetnames:
        if not re.fullmatch(r"[1-6][A-D]", name):
            continue
        ws = wb[name]
        header_row = None
        for i, row in enumerate(ws.iter_rows(max_row=20, values_only=True), start=1):
            if row and row[0] == "No" and row[1] == "English Name":
                header_row = i
                break
        if not header_row:
            continue
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            no = row[0]
            if no is None or no == "":
                continue
            try:
                class_no = int(float(no))
            except (TypeError, ValueError):
                continue
            eng = (row[1] or "").strip() if row[1] else ""
            if not eng:
                continue
            alias = (row[2] or "").strip() if row[2] else ""
            chi = (row[3] or "").strip() if row[3] else ""
            chi = chi.replace("\xa0", " ").strip()
            sex = (row[4] or "").strip() if row[4] else ""
            house = (row[5] or "").strip() if row[5] else ""
            students.append(
                {
                    "class_code": name,
                    "class_no": class_no,
                    "english_name": eng,
                    "alias": alias,
                    "chinese_name": chi,
                    "sex": sex,
                    "house": house,
                    "academic_year": YEAR,
                }
            )
    # dedupe by (class_code, class_no)
    seen = {}
    for s in students:
        key = (s["class_code"], s["class_no"])
        seen[key] = s
    return list(seen.values())


def subject_from_sheet(sheet_name: str, form: str) -> str:
    # strip leading "S2 "/"S3 " etc
    s = re.sub(rf"^{form}\s+", "", sheet_name).strip()
    return s


def parse_group_xls(path: Path, form: str):
    """Return list of group memberships (not yet linked to student_id)."""
    wb = xlrd.open_workbook(str(path))
    groups = []
    for sn in wb.sheet_names():
        if sn in SKIP_SHEETS or "各科分組" in sn:
            continue
        if not sn.startswith(form):
            continue
        # skip ApL with bad teachers later if needed
        sh = wb.sheet_by_name(sn)
        subject = subject_from_sheet(sn, form)
        current_teacher = ""
        header_map = None  # col indices

        for r in range(sh.nrows):
            # detect Teacher :
            for c in range(min(sh.ncols, 8)):
                v = sh.cell_value(r, c)
                if isinstance(v, str) and v.strip().rstrip(":") in ("Teacher", "Teacher "):
                    t = sh.cell_value(r, c + 1) if c + 1 < sh.ncols else ""
                    if isinstance(t, float):
                        # ApL junk
                        current_teacher = ""
                    else:
                        current_teacher = str(t).strip().lower()
                    break
                if isinstance(v, str) and v.strip() == "Teacher :":
                    t = sh.cell_value(r, c + 1) if c + 1 < sh.ncols else ""
                    if isinstance(t, float):
                        current_teacher = ""
                    else:
                        current_teacher = str(t).strip().lower()
                    break

            # detect header
            row_vals = [sh.cell_value(r, c) for c in range(sh.ncols)]
            strs = [str(x).strip() if not isinstance(x, float) else x for x in row_vals]
            if "CL" in strs and "No" in strs and ("English Name" in strs or any(
                isinstance(x, str) and "English" in x for x in strs
            )):
                header_map = {}
                for c, val in enumerate(row_vals):
                    vs = str(val).strip() if val != "" else ""
                    if vs in ("Group", "Subj"):
                        header_map["group"] = c
                    elif vs == "CL":
                        header_map["cl"] = c
                    elif vs == "No":
                        header_map["no"] = c
                    elif vs == "English Name":
                        header_map["eng"] = c
                    elif vs in ("中文姓名",):
                        header_map["chi"] = c
                    elif vs in ("SEX", "Sex"):
                        header_map["sex"] = c
                continue

            if not header_map or "cl" not in header_map or "no" not in header_map:
                continue

            cl = sh.cell_value(r, header_map["cl"])
            no = sh.cell_value(r, header_map["no"])
            if cl == "" or no == "":
                continue
            cl = str(cl).strip()
            if not re.fullmatch(r"[1-6][A-D]", cl):
                continue
            try:
                class_no = int(float(no))
            except (TypeError, ValueError):
                continue

            gname = ""
            if "group" in header_map:
                gv = sh.cell_value(r, header_map["group"])
                gname = str(gv).strip() if gv != "" else ""
            eng = ""
            if "eng" in header_map:
                ev = sh.cell_value(r, header_map["eng"])
                eng = str(ev).strip() if ev != "" else ""

            groups.append(
                {
                    "form": form,
                    "subject": subject,
                    "group_name": gname or subject,
                    "teacher_short_name": current_teacher,
                    "class_code": cl,
                    "class_no": class_no,
                    "english_name": eng,
                    "sheet": sn,
                }
            )
    return groups


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    students = parse_students_xlsx(ATT / "01M511GPB5X3Q8M9PM027JK6MS.xlsx")
    students.sort(key=lambda s: (s["class_code"], s["class_no"]))

    files = [
        ("S6", "01M511GZPP8H3N2EGMVSQAX2JR.xls"),
        ("S5", "01M511H7MWXQN7BPM7QHB5VNSG.xls"),
        ("S4", "01M511J7JSDN7QBK5R8X0SCJZ4.xls"),  # S4 dup skipped
        ("S2", "01M511KDQ9MVXN4CFW975JT1TZ.xls"),
        ("S3", "01M511M9Q7ZGDMY2J70HDEJ1CH.xls"),
    ]
    all_groups = []
    for form, fname in files:
        all_groups.extend(parse_group_xls(ATT / fname, form))

    # dedupe groups by (form, subject, group_name, teacher, class_code, class_no)
    gseen = {}
    for g in all_groups:
        key = (
            g["form"],
            g["subject"],
            g["group_name"],
            g["teacher_short_name"],
            g["class_code"],
            g["class_no"],
        )
        gseen[key] = g
    all_groups = list(gseen.values())

    # index students
    by_class = {(s["class_code"], s["class_no"]): s for s in students}
    unmatched = []
    matched_groups = []
    for g in all_groups:
        key = (g["class_code"], g["class_no"])
        if key not in by_class:
            unmatched.append(g)
            continue
        matched_groups.append(g)

    class_teachers = []
    for cc, shorts in STAFF_CLASS_TEACHERS.items():
        for i, sn in enumerate(shorts):
            class_teachers.append(
                {
                    "class_code": cc,
                    "teacher_short_name": sn,
                    "sort_order": i,
                    "academic_year": YEAR,
                }
            )

    summary = {
        "students": len(students),
        "by_class": {},
        "groups": len(matched_groups),
        "unmatched_groups": len(unmatched),
        "class_teachers": len(class_teachers),
        "subjects_forms": {},
    }
    for s in students:
        summary["by_class"][s["class_code"]] = summary["by_class"].get(s["class_code"], 0) + 1
    for g in matched_groups:
        k = f"{g['form']}|{g['subject']}"
        summary["subjects_forms"][k] = summary["subjects_forms"].get(k, 0) + 1

    (OUT / "students.json").write_text(json.dumps(students, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "groups.json").write_text(json.dumps(matched_groups, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "unmatched_groups.json").write_text(json.dumps(unmatched[:50], ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "class_teachers.json").write_text(json.dumps(class_teachers, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    # Write SQL seed
    lines = []
    lines.append(f"-- Seed students / student_groups / class_teachers  academic_year={YEAR}")
    lines.append("BEGIN;")
    lines.append(f"DELETE FROM public.student_groups WHERE academic_year = '{YEAR}';")
    lines.append(f"DELETE FROM public.class_teachers WHERE academic_year = '{YEAR}';")
    lines.append(f"DELETE FROM public.students WHERE academic_year = '{YEAR}';")
    lines.append("")
    lines.append("-- students")
    batch = []
    for s in students:
        batch.append(
            f"({sql_str(s['class_code'])},{s['class_no']},{sql_str(s['english_name'])},"
            f"{sql_str(s['chinese_name'])},{sql_str(s['sex'])},NULL,{sql_str(YEAR)},"
            f"{sql_str(s['alias'])},{sql_str(s['house'])})"
        )
    # chunk inserts
    cols = "(class_code,class_no,english_name,chinese_name,sex,email,academic_year,alias,house)"
    for i in range(0, len(batch), 80):
        chunk = batch[i : i + 80]
        lines.append(f"INSERT INTO public.students {cols} VALUES\n" + ",\n".join(chunk) + ";")

    lines.append("")
    lines.append("-- class_teachers")
    ct_vals = []
    for ct in class_teachers:
        ct_vals.append(
            f"({sql_str(ct['class_code'])},{sql_str(ct['teacher_short_name'])},"
            f"{ct['sort_order']},{sql_str(YEAR)})"
        )
    lines.append(
        "INSERT INTO public.class_teachers (class_code,teacher_short_name,sort_order,academic_year) VALUES\n"
        + ",\n".join(ct_vals)
        + ";"
    )

    lines.append("")
    lines.append("-- student_groups (join by class_code+class_no+year)")
    gbatch = []
    for g in matched_groups:
        gbatch.append(
            f"((SELECT id FROM public.students WHERE academic_year={sql_str(YEAR)} "
            f"AND class_code={sql_str(g['class_code'])} AND class_no={g['class_no']}),"
            f"{sql_str(g['form'])},{sql_str(g['subject'])},{sql_str(g['group_name'])},"
            f"{sql_str(g['teacher_short_name'])},{sql_str(YEAR)})"
        )
    gcols = "(student_id,form,subject,group_name,teacher_short_name,academic_year)"
    for i in range(0, len(gbatch), 60):
        chunk = gbatch[i : i + 60]
        lines.append(f"INSERT INTO public.student_groups {gcols} VALUES\n" + ",\n".join(chunk) + ";")

    lines.append("COMMIT;")
    seed_path = Path("/workspace/school-admin/db/seed_students.sql")
    seed_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"Wrote {seed_path} ({seed_path.stat().st_size} bytes)")
    print(f"unmatched sample: {unmatched[:5]}")


if __name__ == "__main__":
    main()
